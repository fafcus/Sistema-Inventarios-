from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.styles import Alignment


BASE_DIR = Path(__file__).resolve().parent
PLANTILLA_RELACION_TRANSITO = BASE_DIR / "plantillas" / "RELACION DE TRANSITO.xlsx"


class RelacionTransitoError(Exception):
    """Error controlado al generar una Relación de Tránsito."""


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def _numero(valor):
    try:
        numero = float(valor or 0)
        return int(numero) if numero.is_integer() else numero
    except (TypeError, ValueError):
        return valor if valor is not None else ""


def _celda_escritura(ws, referencia: str):
    """
    Devuelve la celda real donde debe escribirse un valor.

    Excel permite combinar varias celdas, pero openpyxl solo permite escribir
    en la celda superior izquierda del rango combinado. La plantilla de
    Relación de Tránsito utiliza varios merges, por lo que no debemos asumir
    que referencias como B6 o B7 son directamente escribibles.
    """
    celda = ws[referencia]

    for rango in ws.merged_cells.ranges:
        if celda.coordinate in rango:
            return ws.cell(rango.min_row, rango.min_col)

    return celda


def _asignar_valor(ws, referencia: str, valor) -> None:
    """Escribe un valor respetando automáticamente las celdas combinadas."""
    _celda_escritura(ws, referencia).value = valor


def _asignar_valor_fila(ws, fila: int, columna: int, valor) -> None:
    """Versión por fila/columna que también soporta rangos combinados."""
    referencia = ws.cell(fila, columna).coordinate
    _asignar_valor(ws, referencia, valor)


def _copiar_estilo(origen, destino) -> None:
    """Copia el formato de una celda de la plantilla sin copiar su valor."""
    if origen.has_style:
        destino._style = copy(origen._style)
    if origen.number_format:
        destino.number_format = origen.number_format
    if origen.alignment:
        destino.alignment = copy(origen.alignment)
    if origen.protection:
        destino.protection = copy(origen.protection)
    if origen.font:
        destino.font = copy(origen.font)
    if origen.fill:
        destino.fill = copy(origen.fill)
    if origen.border:
        destino.border = copy(origen.border)


def _preparar_fila(ws, fila: int, fila_modelo: int = 12) -> None:
    """Replica el formato de la primera fila de carga de la plantilla."""
    for columna in range(1, 8):
        origen = _celda_escritura(ws, ws.cell(fila_modelo, columna).coordinate)
        destino = _celda_escritura(ws, ws.cell(fila, columna).coordinate)
        _copiar_estilo(origen, destino)
    ws.row_dimensions[fila].height = ws.row_dimensions[fila_modelo].height


def _observaciones_con_ubicacion(material: Mapping) -> str:
    """
    La ubicación se incorpora siempre a Observaciones.

    Si el material ya tiene observaciones, se conservan y se agrega la
    ubicación después. Esto evita perder información existente.
    """
    ubicacion = _texto(material.get("ubicacion"))
    observaciones = _texto(material.get("observaciones"))

    if ubicacion:
        texto_ubicacion = f"Ubicación: {ubicacion}"
        if observaciones:
            return f"{observaciones} | {texto_ubicacion}"
        return texto_ubicacion

    return observaciones


def generar_relacion_transito(
    materiales: Iterable[Mapping],
    salida: str | Path,
    tipo_movimiento: str = "",
    destino: str = "",
    transporte: str = "",
    fecha: datetime | str | None = None,
    plantilla: str | Path | None = None,
) -> Path:
    """
    Genera una Relación de Tránsito usando exactamente la plantilla Excel.

    Cada elemento de ``materiales`` puede contener:
      codigo, material, cantidad, unidad, ubicacion, observaciones,
      marca, numero_serie y numero_parte.

    Mapeo a la plantilla:
      A = N°
      B = CANTIDAD
      C = PARTES (material)
      D = MARCA
      E = NUMERO DE SERIE
      F = NUMERO DE PARTE (código)
      G = OBSERVACIONES / UBICACIÓN
    """
    plantilla = Path(plantilla or PLANTILLA_RELACION_TRANSITO)
    salida = Path(salida)

    if not plantilla.exists():
        raise RelacionTransitoError(
            f"No se encontró la plantilla de Relación de Tránsito: {plantilla}"
        )

    materiales = list(materiales or [])
    if not materiales:
        raise RelacionTransitoError("No hay materiales para generar la relación de tránsito.")

    try:
        wb = load_workbook(plantilla)
    except Exception as error:
        raise RelacionTransitoError(f"No se pudo abrir la plantilla Excel: {error}") from error

    if "Hoja1" not in wb.sheetnames:
        raise RelacionTransitoError("La plantilla no contiene la hoja 'Hoja1'.")

    ws = wb["Hoja1"]

    # Encabezado de la relación.
    _asignar_valor(ws, "G11", "OBSERVACIONES / UBICACIÓN")
    _copiar_estilo(_celda_escritura(ws, "F11"), _celda_escritura(ws, "G11"))
    _celda_escritura(ws, "G11").alignment = copy(_celda_escritura(ws, "F11").alignment)

    if fecha is None:
        fecha_texto = datetime.now().strftime("%d/%m/%Y")
    elif isinstance(fecha, datetime):
        fecha_texto = fecha.strftime("%d/%m/%Y")
    else:
        fecha_texto = _texto(fecha)

    _asignar_valor(ws, "G2", fecha_texto)
    _asignar_valor(ws, "B6", _texto(tipo_movimiento))
    _asignar_valor(ws, "B7", _texto(destino))
    _asignar_valor(ws, "B8", _texto(transporte))

    fila_inicial = 12
    fila_modelo = 12
    cantidad_filas_disponibles = max(0, ws.max_row - fila_inicial + 1)

    # La plantilla trae filas numeradas. Si hay más materiales, agregamos filas
    # manteniendo el estilo de la fila modelo.
    filas_necesarias = len(materiales)
    if filas_necesarias > cantidad_filas_disponibles:
        filas_extra = filas_necesarias - cantidad_filas_disponibles
        ws.insert_rows(ws.max_row + 1, amount=filas_extra)
        for fila in range(ws.max_row - filas_extra + 1, ws.max_row + 1):
            _preparar_fila(ws, fila, fila_modelo)

    for indice, material in enumerate(materiales, start=1):
        fila = fila_inicial + indice - 1

        if fila != fila_modelo:
            _preparar_fila(ws, fila, fila_modelo)

        codigo = _texto(material.get("codigo"))
        nombre = _texto(material.get("material"))
        marca = _texto(material.get("marca"))
        numero_serie = _texto(material.get("numero_serie"))
        numero_parte = _texto(material.get("numero_parte")) or codigo
        observaciones = _observaciones_con_ubicacion(material)

        _asignar_valor_fila(ws, fila, 1, indice)
        _asignar_valor_fila(ws, fila, 2, _numero(material.get("cantidad")))
        _asignar_valor_fila(ws, fila, 3, nombre)
        _asignar_valor_fila(ws, fila, 4, marca)
        _asignar_valor_fila(ws, fila, 5, numero_serie)
        _asignar_valor_fila(ws, fila, 6, numero_parte)
        _asignar_valor_fila(ws, fila, 7, observaciones)

        _celda_escritura(ws, ws.cell(fila, 7).coordinate).alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )

    # Limpiar filas restantes de la plantilla para evitar números sobrantes.
    ultima_fila = fila_inicial + filas_necesarias - 1
    for fila in range(ultima_fila + 1, ws.max_row + 1):
        for columna in range(1, 8):
            _asignar_valor_fila(ws, fila, columna, None)

    ws.freeze_panes = "A12"
    ws.sheet_view.showGridLines = False

    # Aseguramos que Observaciones tenga espacio suficiente sin modificar el
    # resto de la composición visual de la plantilla.
    if ws.column_dimensions["G"].width < 32:
        ws.column_dimensions["G"].width = 32

    salida.parent.mkdir(parents=True, exist_ok=True)

    try:
        wb.save(salida)
    except Exception as error:
        raise RelacionTransitoError(f"No se pudo guardar la Relación de Tránsito: {error}") from error

    return salida


def generar_desde_inventario(
    materiales: Iterable[Mapping],
    carpeta_salida: str | Path,
    tipo_movimiento: str = "",
    destino: str = "",
    transporte: str = "",
    fecha: datetime | str | None = None,
) -> Path:
    """Atajo para guardar la relación con un nombre fechado."""
    carpeta_salida = Path(carpeta_salida)
    marca_fecha = datetime.now().strftime("%Y%m%d_%H%M%S")
    salida = carpeta_salida / f"RELACION DE TRANSITO_{marca_fecha}.xlsx"
    return generar_relacion_transito(
        materiales=materiales,
        salida=salida,
        tipo_movimiento=tipo_movimiento,
        destino=destino,
        transporte=transporte,
        fecha=fecha,
    )
