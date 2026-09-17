import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
from pathlib import Path
from copy import deepcopy
from datetime import datetime
import threading
import time
import traceback

from supabase_db import (
    probar_conexion,
    obtener_materiales,
    obtener_material,
    crear_material,
    actualizar_material,
    actualizar_cantidad_documento_item,
    obtener_movimientos,
    obtener_documentos,
    obtener_items_documento,
    crear_item_documento,
    buscar_material,
    eliminar_material_del_documento,
)
import importar_word
from relacion_transito import RelacionTransitoError, generar_relacion_transito


# ============================================================
# CONFIGURACIÓN
# ============================================================

CARPETA_DOCUMENTOS = Path(__file__).resolve().parent / "documentos"

NOMBRE_APP = "INVENTARIO MATERIAL NAVAL"

COLOR_STOCK = "🟢 HAY STOCK"
COLOR_SIN_STOCK = "🔴 SIN STOCK"

INTERVALO_MONITOR = 5000
INTERVALO_SINCRONIZACION = 3000

COLOR_FONDO = "#eef3f8"
COLOR_PANEL = "#ffffff"
COLOR_AZUL_OSCURO = "#12304a"
COLOR_AZUL = "#1f5d87"
COLOR_AZUL_CLARO = "#dceaf5"
COLOR_BORDE = "#c7d5e0"
COLOR_TEXTO = "#183247"
COLOR_TEXTO_SECUNDARIO = "#506575"
COLOR_STOCK_FONDO = "#e7f6e9"
COLOR_SIN_STOCK_FONDO = "#fdeaea"


# ============================================================
# VARIABLES GLOBALES
# ============================================================

root = None

tabla = None
tabla_movimientos = None
entrada_busqueda = None

lbl_total_materiales = None
lbl_con_stock = None
lbl_sin_stock = None
lbl_cantidad_total = None

lbl_estado = None
lbl_progreso = None

# Inventario seleccionado
inventario_seleccionado = None
inventario_seleccionado_id = None

marco_contenido = None
marco_selector = None

importacion_en_curso = False
lock_importacion = threading.Lock()

estado_archivos_word = {}

cache_materiales = []
cache_documentos = []
cache_movimientos = []

cache_lock = threading.Lock()

actualizacion_en_curso = False
lock_actualizacion = threading.Lock()

sincronizacion_en_curso = False
lock_sincronizacion = threading.Lock()

firma_datos_sincronizados = None

pausar_sincronizacion = False

# Estado de la carga inicial de documentos/materiales.
datos_iniciales_cargados = False
error_carga_inicial = None


# ============================================================
# UTILIDADES
# ============================================================

def limpiar_texto(valor):
    return "" if valor is None else str(valor).strip()


def formatear_numero(valor):
    try:
        return f"{float(valor or 0):g}"
    except Exception:
        return str(valor or 0)


# ============================================================
# CACHE
# ============================================================

def invalidar_cache():
    global cache_materiales
    global cache_documentos
    global cache_movimientos

    with cache_lock:
        cache_materiales = []
        cache_documentos = []
        cache_movimientos = []


def cargar_datos_supabase():
    global cache_materiales
    global cache_documentos
    global cache_movimientos

    materiales = obtener_materiales() or []
    documentos = obtener_documentos() or []
    movimientos = obtener_movimientos(100) or []

    with cache_lock:
        cache_materiales = materiales
        cache_documentos = documentos
        cache_movimientos = movimientos

    return materiales, documentos, movimientos


def obtener_materiales_cache():
    with cache_lock:
        return list(cache_materiales)


def obtener_documentos_cache():
    with cache_lock:
        return list(cache_documentos)


def obtener_movimientos_cache():
    with cache_lock:
        return list(cache_movimientos)


def generar_firma_sincronizacion(
    materiales,
    documentos,
    movimientos
):
    return (
        tuple(
            (
                m.get("id"),
                m.get("codigo"),
                m.get("material"),
                m.get("unidad"),
                m.get("categoria"),
                m.get("ubicacion"),
                m.get("observaciones"),
            )
            for m in materiales
        ),
        tuple(
            (
                d.get("id"),
                d.get("nombre"),
                d.get("ruta"),
                d.get("fecha_modificacion_archivo"),
            )
            for d in documentos
        ),
        tuple(
            (
                m.get("id"),
                m.get("material_id"),
                m.get("tipo"),
                m.get("cantidad"),
                m.get("stock_anterior"),
                m.get("stock_nuevo"),
                m.get("fecha"),
                m.get("usuario"),
                m.get("archivo_origen"),
                m.get("observaciones"),
            )
            for m in movimientos
        ),
    )


# ============================================================
# ESTADÍSTICAS
# ============================================================

def actualizar_estadisticas(materiales):
    total = len(materiales)

    con = 0
    sin = 0
    total_cantidad = 0

    for m in materiales:
        try:
            cantidad = float(m.get("cantidad", 0) or 0)
        except Exception:
            cantidad = 0

        total_cantidad += cantidad

        if cantidad > 0:
            con += 1
        else:
            sin += 1

    if lbl_total_materiales:
        try:
            lbl_total_materiales.config(
                text=f"Materiales: {total}"
            )
        except Exception:
            pass

    if lbl_con_stock:
        try:
            lbl_con_stock.config(
                text=f"Con stock: {con}"
            )
        except Exception:
            pass

    if lbl_sin_stock:
        try:
            lbl_sin_stock.config(
                text=f"Sin stock: {sin}"
            )
        except Exception:
            pass

    if lbl_cantidad_total:
        try:
            lbl_cantidad_total.config(
                text=f"Cantidad total: {total_cantidad:g}"
            )
        except Exception:
            pass


# ============================================================
# DOCUMENTO / INVENTARIO ACTUAL
# ============================================================

def obtener_documento_actual():
    """
    Devuelve EXACTAMENTE el documento seleccionado.

    Primero busca por ID.
    El nombre queda solamente como respaldo.
    """

    global inventario_seleccionado_id
    global inventario_seleccionado

    documentos = obtener_documentos_cache()

    # --------------------------------------------------------
    # 1. BUSCAR POR ID
    # --------------------------------------------------------

    if inventario_seleccionado_id is not None:

        for documento in documentos:

            if documento.get("id") == inventario_seleccionado_id:
                return documento

    # --------------------------------------------------------
    # 2. RESPALDO: BUSCAR POR NOMBRE
    # --------------------------------------------------------

    if inventario_seleccionado:

        for documento in documentos:

            if documento.get("nombre") == inventario_seleccionado:
                return documento

    return None


def obtener_documento_id_actual():

    documento = obtener_documento_actual()

    if not documento:
        return None

    return documento.get("id")


# ============================================================
# MATERIALES DEL INVENTARIO ACTUAL
# ============================================================

def obtener_materiales_seleccionados():

    documento_id = obtener_documento_id_actual()

    if not documento_id:
        return []

    try:
        items = obtener_items_documento(documento_id) or []

    except Exception as error:
        print(
            "Error obteniendo documento_items:",
            error
        )
        return []

    materiales = obtener_materiales_cache()

    por_id = {
        m.get("id"): m
        for m in materiales
        if m.get("id") is not None
    }

    resultado = []

    for item in items:

        mid = item.get("material_id")

        material = por_id.get(mid)

        if not material:

            try:
                material = obtener_material(mid)

            except Exception:
                material = None

        if not material:
            continue

        copia = dict(material)

        copia["cantidad"] = item.get(
            "cantidad",
            0
        )

        copia["_documento_item_id"] = item.get("id")

        for campo in (
            "codigo",
            "material",
            "unidad",
            "categoria",
            "ubicacion",
            "observaciones",
        ):

            if item.get(campo) is not None:
                copia[campo] = item.get(campo)

        resultado.append(copia)

    return resultado


# ============================================================
# TABLA PRINCIPAL
# ============================================================

def actualizar_tabla():

    if tabla is None:
        return

    try:
        for x in tabla.get_children():
            tabla.delete(x)
    except Exception:
        return

    materiales = obtener_materiales_seleccionados()

    texto = ""

    if entrada_busqueda:

        try:
            texto = limpiar_texto(
                entrada_busqueda.get()
            ).lower()
        except Exception:
            texto = ""

    filtrados = []

    for m in materiales:

        if texto:

            combinado = " ".join(
                limpiar_texto(
                    m.get(c)
                )
                for c in (
                    "codigo",
                    "material",
                    "categoria",
                    "ubicacion",
                    "observaciones",
                )
            ).lower()

            if texto not in combinado:
                continue

        filtrados.append(m)

    for m in filtrados:

        try:
            cantidad = float(
                m.get("cantidad", 0) or 0
            )

        except Exception:
            cantidad = 0

        if cantidad > 0:
            estado = COLOR_STOCK
            tag = "stock"

        else:
            estado = COLOR_SIN_STOCK
            tag = "sin_stock"

        try:

            tabla.insert(
                "",
                "end",
                iid=str(m.get("id")),
                values=(
                    m.get("codigo") or "",
                    m.get("material") or "",
                    formatear_numero(cantidad),
                    m.get("unidad") or "",
                    m.get("categoria") or "",
                    m.get("ubicacion") or "",
                    m.get("observaciones") or "",
                    estado,
                ),
                tags=(tag,),
            )

        except Exception:
            pass

    actualizar_estadisticas(
        filtrados
    )


def buscar(event=None):
    actualizar_tabla()


# ============================================================
# MATERIAL SELECCIONADO
# ============================================================

def obtener_material_seleccionado():

    if tabla is None:
        return None

    try:
        sel = tabla.selection()

    except Exception:
        return None

    if not sel:

        messagebox.showwarning(
            "Selección",
            "Seleccioná un material primero.",
            parent=root,
        )

        return None

    try:
        mid = int(sel[0])

    except Exception:
        return None

    for m in obtener_materiales_seleccionados():

        if m.get("id") == mid:
            return m

    try:
        return obtener_material(mid)

    except Exception:
        return None


# ============================================================
# ITEM DEL DOCUMENTO
# ============================================================

def obtener_item_documento(material_id):

    documento_id = obtener_documento_id_actual()

    if not documento_id:
        return None

    try:
        items = obtener_items_documento(
            documento_id
        ) or []

    except Exception:
        return None

    for item in items:

        try:

            if int(
                item.get("material_id", -1)
            ) == int(material_id):

                return item

        except Exception:
            pass

    return None


# ============================================================
# ACTUALIZAR STOCK
# ============================================================

def actualizar_stock_documento(
    material_id,
    nueva_cantidad,
    tipo_movimiento=None,
    usuario="APP",
    observaciones=None,
):

    item = obtener_item_documento(
        material_id
    )

    if item is None:

        raise Exception(
            "El material no pertenece al inventario seleccionado."
        )

    return actualizar_cantidad_documento_item(
        item["id"],
        nueva_cantidad,
        usuario=usuario,
        archivo_origen=inventario_seleccionado,
        observaciones=observaciones,
        documento_id=obtener_documento_id_actual(),
    )


# ============================================================
# AGREGAR STOCK
# ============================================================

def agregar_stock():

    material = obtener_material_seleccionado()

    if not material:
        return

    mid = material["id"]

    nombre = material.get("material") or ""
    unidad = material.get("unidad") or ""
    codigo = material.get("codigo") or "-"

    item = obtener_item_documento(mid)

    if item is None:

        messagebox.showerror(
            "Error",
            "El material no pertenece al inventario seleccionado.",
            parent=root,
        )

        return

    stock = float(
        item.get("cantidad", 0) or 0
    )

    cantidad = simpledialog.askfloat(
        "Agregar stock",
        f"Inventario: {inventario_seleccionado}\n"
        f"Material: {nombre}\n"
        f"Código: {codigo}\n"
        f"Stock actual: {formatear_numero(stock)} {unidad}\n\n"
        f"Cantidad a agregar:",
        minvalue=0.0001,
        parent=root,
    )

    if cantidad is None or cantidad <= 0:
        return

    try:

        nuevo = actualizar_stock_documento(
            mid,
            stock + cantidad,
            "ENTRADA",
            "APP",
            "Ingreso manual en "
            + inventario_seleccionado,
        )

        invalidar_cache()
        actualizar_todo(True)

        valor = (
            nuevo.get(
                "cantidad",
                stock + cantidad,
            )
            if isinstance(nuevo, dict)
            else stock + cantidad
        )

        messagebox.showinfo(
            "Stock actualizado",
            f"Material: {nombre}\n\n"
            f"Stock anterior: {formatear_numero(stock)} {unidad}\n"
            f"Cantidad agregada: {formatear_numero(cantidad)} {unidad}\n"
            f"Stock nuevo: {formatear_numero(valor)} {unidad}",
            parent=root,
        )

    except Exception as error:

        traceback.print_exc()

        messagebox.showerror(
            "Error",
            f"No se pudo agregar stock:\n\n{error}",
            parent=root,
        )


# ============================================================
# RETIRAR STOCK
# ============================================================

def retirar_stock():

    material = obtener_material_seleccionado()

    if not material:
        return

    mid = material["id"]

    nombre = material.get("material") or ""
    unidad = material.get("unidad") or ""
    codigo = material.get("codigo") or "-"

    item = obtener_item_documento(mid)

    if item is None:

        messagebox.showerror(
            "Error",
            "El material no pertenece al inventario seleccionado.",
            parent=root,
        )

        return

    stock = float(
        item.get("cantidad", 0) or 0
    )

    cantidad = simpledialog.askfloat(
        "Retirar stock",
        f"Inventario: {inventario_seleccionado}\n"
        f"Material: {nombre}\n"
        f"Código: {codigo}\n"
        f"Stock actual: {formatear_numero(stock)} {unidad}\n\n"
        f"Cantidad a retirar:",
        minvalue=0.0001,
        parent=root,
    )

    if cantidad is None or cantidad <= 0:
        return

    if cantidad > stock:

        messagebox.showwarning(
            "Stock insuficiente",
            f"No podés retirar "
            f"{formatear_numero(cantidad)} {unidad}.\n\n"
            f"Stock disponible: "
            f"{formatear_numero(stock)} {unidad}",
            parent=root,
        )

        return

    try:

        nuevo = actualizar_stock_documento(
            mid,
            stock - cantidad,
            "SALIDA",
            "APP",
            "Retiro manual en "
            + inventario_seleccionado,
        )

        invalidar_cache()
        actualizar_todo(True)

        valor = (
            nuevo.get(
                "cantidad",
                stock - cantidad,
            )
            if isinstance(nuevo, dict)
            else stock - cantidad
        )

        messagebox.showinfo(
            "Retiro realizado",
            f"Material: {nombre}\n\n"
            f"Stock anterior: {formatear_numero(stock)} {unidad}\n"
            f"Cantidad retirada: {formatear_numero(cantidad)} {unidad}\n"
            f"Stock nuevo: {formatear_numero(valor)} {unidad}",
            parent=root,
        )

    except Exception as error:

        traceback.print_exc()

        messagebox.showerror(
            "Error",
            f"No se pudo retirar stock:\n\n{error}",
            parent=root,
        )


# ============================================================
# AGREGAR FILA AL WORD
# ============================================================

def _agregar_fila_word(documento, datos):

    try:
        from docx import Document

    except Exception as error:

        raise Exception(
            "Falta instalar python-docx. "
            "Ejecutá: pip install python-docx"
        ) from error

    ruta = Path(
        str(documento.get("ruta") or "")
    )

    if not ruta.exists():

        raise Exception(
            f"No se encontró el archivo Word:\n{ruta}"
        )

    doc = Document(str(ruta))

    tabla_objetivo = None
    mapa = None

    for tabla_word in doc.tables:

        if not tabla_word.rows:
            continue

        encabezados = [
            limpiar_texto(c.text).lower()
            for c in tabla_word.rows[0].cells
        ]

        mapa_tmp = {}

        for i, encabezado in enumerate(encabezados):

            if encabezado in (
                "codigo",
                "código",
            ):
                mapa_tmp["codigo"] = i

            elif encabezado == "material":
                mapa_tmp["material"] = i

            elif encabezado in (
                "cantidad",
                "cant",
                "cant.",
                "stock",
                "existencia",
                "existencias",
            ):
                mapa_tmp["cantidad"] = i

            elif encabezado == "unidad":
                mapa_tmp["unidad"] = i

            elif encabezado in (
                "categoria",
                "categoría",
            ):
                mapa_tmp["categoria"] = i

            elif encabezado in (
                "ubicacion",
                "ubicación",
            ):
                mapa_tmp["ubicacion"] = i

            elif encabezado in (
                "observaciones",
                "obs",
            ):
                mapa_tmp["observaciones"] = i

        if (
            "material" in mapa_tmp
            and "cantidad" in mapa_tmp
        ):

            tabla_objetivo = tabla_word
            mapa = mapa_tmp
            break

    if tabla_objetivo is None:

        raise Exception(
            "No encontré en el Word una tabla "
            "con las columnas Material y Cantidad."
        )

    nueva_fila = tabla_objetivo.add_row()

    if len(tabla_objetivo.rows) >= 2:

        try:

            fila_modelo = tabla_objetivo.rows[-2]

            for origen, destino in zip(
                fila_modelo.cells,
                nueva_fila.cells,
            ):

                destino._tc.get_or_add_tcPr()

                for hijo in list(
                    origen._tc.tcPr
                ):

                    destino._tc.tcPr.append(
                        deepcopy(hijo)
                    )

        except Exception:
            pass

    valores = {
        "codigo": datos.get("codigo") or "",
        "material": datos.get("material") or "",
        "cantidad": formatear_numero(
            datos.get("cantidad", 0)
        ),
        "unidad": datos.get("unidad") or "",
        "categoria": datos.get("categoria") or "",
        "ubicacion": datos.get("ubicacion") or "",
        "observaciones": datos.get("observaciones") or "",
    }

    for campo, indice in mapa.items():

        if indice < len(
            nueva_fila.cells
        ):

            nueva_fila.cells[
                indice
            ].text = valores.get(
                campo,
                "",
            )

    doc.save(str(ruta))

    return ruta


# ============================================================
# ELIMINAR FILA DEL WORD
# ============================================================

def _eliminar_fila_word(documento, material):

    try:
        from docx import Document

    except Exception as error:

        raise Exception(
            "Falta instalar python-docx. "
            "Ejecutá: pip install python-docx"
        ) from error

    ruta = Path(
        str(documento.get("ruta") or "")
    )

    if not ruta.exists():

        raise Exception(
            f"No se encontró el archivo Word:\n{ruta}"
        )

    doc = Document(str(ruta))

    codigo = limpiar_texto(
        material.get("codigo")
    )

    nombre = limpiar_texto(
        material.get("material")
    )

    fila_encontrada = False

    for tabla_word in doc.tables:

        if not tabla_word.rows:
            continue

        encabezados = [
            limpiar_texto(c.text).lower()
            for c in tabla_word.rows[0].cells
        ]

        indice_codigo = None
        indice_material = None

        for i, encabezado in enumerate(
            encabezados
        ):

            if encabezado in (
                "codigo",
                "código",
            ):

                indice_codigo = i

            elif encabezado == "material":

                indice_material = i

        if (
            indice_codigo is None
            and indice_material is None
        ):
            continue

        if codigo and indice_codigo is not None:

            for fila in list(
                tabla_word.rows[1:]
            ):

                valores = [
                    limpiar_texto(
                        celda.text
                    )
                    for celda in fila.cells
                ]

                if (
                    indice_codigo < len(valores)
                    and valores[indice_codigo] == codigo
                ):

                    tr = fila._tr

                    tr.getparent().remove(tr)

                    fila_encontrada = True
                    break

        elif (
            not codigo
            and nombre
            and indice_material is not None
        ):

            for fila in list(
                tabla_word.rows[1:]
            ):

                valores = [
                    limpiar_texto(
                        celda.text
                    )
                    for celda in fila.cells
                ]

                if (
                    indice_material < len(valores)
                    and valores[indice_material] == nombre
                ):

                    tr = fila._tr

                    tr.getparent().remove(tr)

                    fila_encontrada = True
                    break

        if fila_encontrada:
            break

    if not fila_encontrada:

        raise Exception(
            "No se encontró el material seleccionado "
            "en el archivo Word."
        )

    doc.save(str(ruta))

    return ruta


# ============================================================
# NUEVO MATERIAL
# ============================================================

def nuevo_material():

    if not inventario_seleccionado:
        return

    ventana = tk.Toplevel(root)

    ventana.title(
        "Nuevo material"
    )

    ventana.geometry(
        "520x520"
    )

    ventana.transient(root)
    ventana.grab_set()

    campos = [
        ("Código", "codigo"),
        ("Material", "material"),
        ("Cantidad", "cantidad"),
        ("Unidad", "unidad"),
        ("Categoría", "categoria"),
        ("Ubicación", "ubicacion"),
        ("Observaciones", "observaciones"),
    ]

    entradas = {}

    frame = ttk.Frame(
        ventana,
        padding=20,
    )

    frame.pack(
        fill="both",
        expand=True,
    )

    for fila, (texto, clave) in enumerate(campos):

        ttk.Label(
            frame,
            text=texto,
        ).grid(
            row=fila,
            column=0,
            sticky="w",
            padx=5,
            pady=7,
        )

        e = ttk.Entry(frame)

        e.grid(
            row=fila,
            column=1,
            sticky="ew",
            padx=5,
            pady=7,
        )

        entradas[clave] = e

    frame.columnconfigure(
        1,
        weight=1,
    )

    def guardar():

        global importacion_en_curso

        nombre = (
            entradas["material"]
            .get()
            .strip()
        )

        if not nombre:

            messagebox.showwarning(
                "Datos",
                "El nombre del material es obligatorio.",
                parent=ventana,
            )

            return

        try:

            cantidad = float(
                entradas["cantidad"]
                .get()
                .strip()
                .replace(",", ".")
                or 0
            )

        except ValueError:

            messagebox.showerror(
                "Cantidad",
                "La cantidad debe ser numérica.",
                parent=ventana,
            )

            return

        if cantidad < 0:

            messagebox.showerror(
                "Cantidad",
                "La cantidad no puede ser negativa.",
                parent=ventana,
            )

            return

        datos = {
            k: entradas[k].get().strip() or None
            for k in (
                "codigo",
                "unidad",
                "categoria",
                "ubicacion",
                "observaciones",
            )
        }

        datos["material"] = nombre
        datos["cantidad"] = cantidad

        try:

            if buscar_material(
                datos["codigo"],
                nombre,
                datos["categoria"],
                datos["ubicacion"],
            ):

                messagebox.showwarning(
                    "Material existente",
                    "Ese material ya existe en la base de datos.",
                    parent=ventana,
                )

                return

            documento = obtener_documento_actual()

            if not documento:

                raise Exception(
                    "No se encontró el documento seleccionado."
                )

            with lock_importacion:

                if importacion_en_curso:

                    messagebox.showinfo(
                        "Importación",
                        "Hay una importación de Word en curso. "
                        "Esperá a que termine.",
                        parent=ventana,
                    )

                    return

                importacion_en_curso = True

            try:

                _agregar_fila_word(
                    documento,
                    datos,
                )

                nuevo = crear_material(
                    datos["codigo"],
                    nombre,
                    0,
                    datos["unidad"],
                    datos["categoria"],
                    datos["ubicacion"],
                    datos["observaciones"],
                    inventario_seleccionado,
                )

                if not nuevo:

                    raise Exception(
                        "No se pudo crear el material "
                        "en la base de datos."
                    )

                if not crear_item_documento(
                    documento["id"],
                    nuevo["id"],
                    cantidad,
                    unidad=datos["unidad"],
                    codigo=datos["codigo"],
                    material=nombre,
                    categoria=datos["categoria"],
                    ubicacion=datos["ubicacion"],
                    observaciones=datos["observaciones"],
                ):

                    raise Exception(
                        "No se pudo crear el item del inventario."
                    )

                try:

                    from supabase_db import actualizar_documento

                    actualizar_documento(
                        documento["id"],
                        Path(
                            str(
                                documento["ruta"]
                            )
                        ).stat().st_mtime,
                    )

                except Exception:
                    pass

            finally:

                importacion_en_curso = False

            invalidar_cache()

            ventana.destroy()

            actualizar_todo(True)

            messagebox.showinfo(
                "Material creado",
                f"Se agregó correctamente a "
                f"{inventario_seleccionado}:\n\n"
                f"{nombre}\n"
                f"Cantidad: "
                f"{formatear_numero(cantidad)}",
                parent=root,
            )

        except Exception as error:

            traceback.print_exc()

            messagebox.showerror(
                "Error",
                f"No se pudo crear el material:\n\n{error}",
                parent=ventana,
            )

    ttk.Button(
        frame,
        text="Guardar",
        command=guardar,
    ).grid(
        row=len(campos),
        column=0,
        columnspan=2,
        pady=20,
    )


# ============================================================
# EDITAR MATERIAL
# ============================================================

def editar_material():

    material = obtener_material_seleccionado()

    if not material:
        return

    ventana = tk.Toplevel(root)

    ventana.title(
        "Editar material"
    )

    ventana.geometry(
        "520x500"
    )

    ventana.transient(root)
    ventana.grab_set()

    campos = [
        ("Código", "codigo"),
        ("Material", "material"),
        ("Unidad", "unidad"),
        ("Categoría", "categoria"),
        ("Ubicación", "ubicacion"),
        ("Observaciones", "observaciones"),
    ]

    entradas = {}

    frame = ttk.Frame(
        ventana,
        padding=20,
    )

    frame.pack(
        fill="both",
        expand=True,
    )

    for fila, (texto, clave) in enumerate(campos):

        ttk.Label(
            frame,
            text=texto,
        ).grid(
            row=fila,
            column=0,
            sticky="w",
            padx=5,
            pady=7,
        )

        e = ttk.Entry(frame)

        e.grid(
            row=fila,
            column=1,
            sticky="ew",
            padx=5,
            pady=7,
        )

        e.insert(
            0,
            material.get(clave) or "",
        )

        entradas[clave] = e

    frame.columnconfigure(
        1,
        weight=1,
    )

    def guardar():

        nombre = (
            entradas["material"]
            .get()
            .strip()
        )

        if not nombre:

            messagebox.showwarning(
                "Datos",
                "El material no puede quedar vacío.",
                parent=ventana,
            )

            return

        try:

            if not actualizar_material(
                material["id"],
                codigo=(
                    entradas["codigo"]
                    .get()
                    .strip()
                    or None
                ),
                material=nombre,
                unidad=(
                    entradas["unidad"]
                    .get()
                    .strip()
                    or None
                ),
                categoria=(
                    entradas["categoria"]
                    .get()
                    .strip()
                    or None
                ),
                ubicacion=(
                    entradas["ubicacion"]
                    .get()
                    .strip()
                    or None
                ),
                observaciones=(
                    entradas["observaciones"]
                    .get()
                    .strip()
                    or None
                ),
            ):

                raise Exception(
                    "No se pudo actualizar."
                )

            invalidar_cache()

            ventana.destroy()

            actualizar_todo(True)

        except Exception as error:

            traceback.print_exc()

            messagebox.showerror(
                "Error",
                f"No se pudo modificar:\n\n{error}",
                parent=ventana,
            )

    ttk.Button(
        frame,
        text="Guardar cambios",
        command=guardar,
    ).grid(
        row=len(campos),
        column=0,
        columnspan=2,
        pady=20,
    )


# ============================================================
# ELIMINAR MATERIAL
# ============================================================

def eliminar_material():

    global importacion_en_curso
    global estado_archivos_word

    material = obtener_material_seleccionado()

    if not material:
        return

    documento = obtener_documento_actual()

    if not documento:

        messagebox.showerror(
            "Error",
            "No se encontró el inventario seleccionado.",
            parent=root,
        )

        return

    material_id = material.get("id")

    if material_id is None:

        messagebox.showerror(
            "Error",
            "El material seleccionado no tiene un ID válido.",
            parent=root,
        )

        return

    nombre = (
        material.get("material")
        or ""
    )

    codigo = (
        material.get("codigo")
        or "-"
    )

    unidad = (
        material.get("unidad")
        or ""
    )

    try:

        cantidad = float(
            material.get(
                "cantidad",
                0,
            )
            or 0
        )

    except Exception:

        cantidad = 0

    confirmar = messagebox.askyesno(
        "Eliminar material",

        f"¿Eliminar este material del inventario "
        f"y del archivo Word?\n\n"

        f"Código: {codigo}\n"
        f"Material: {nombre}\n"
        f"Cantidad: "
        f"{formatear_numero(cantidad)} "
        f"{unidad}\n\n"

        f"El material será eliminado del inventario "
        f"seleccionado y del Word.\n\n"

        f"El historial de movimientos se conservará.",

        parent=root,
    )

    if not confirmar:
        return

    with lock_importacion:

        if importacion_en_curso:

            messagebox.showinfo(
                "Importación",
                "Hay una importación de Word en curso. "
                "Esperá a que termine.",
                parent=root,
            )

            return

        importacion_en_curso = True

    try:

        ruta = _eliminar_fila_word(
            documento,
            material,
        )

        resultado = (
            eliminar_material_del_documento(
                documento["id"],
                material_id,
            )
        )

        try:

            from supabase_db import (
                actualizar_documento
            )

            if ruta.exists():

                actualizar_documento(
                    documento["id"],
                    ruta.stat().st_mtime,
                )

        except Exception:

            pass

        try:

            estado_archivos_word[
                str(ruta.resolve())
            ] = ruta.stat().st_mtime

        except Exception:

            pass

        invalidar_cache()

        actualizar_todo(True)

        if (
            isinstance(resultado, dict)
            and resultado.get(
                "material_eliminado"
            )
        ):

            mensaje = (
                f"Se eliminó correctamente:\n\n"
                f"{nombre}\n\n"
                f"✓ Eliminado del Word\n"
                f"✓ Eliminado del inventario\n"
                f"✓ Eliminado de la base general\n"
                f"✓ Historial de movimientos conservado"
            )

        else:

            mensaje = (
                f"Se eliminó correctamente:\n\n"
                f"{nombre}\n\n"
                f"✓ Eliminado del Word\n"
                f"✓ Eliminado del inventario seleccionado\n"
                f"✓ El material continúa en otros inventarios\n"
                f"✓ Historial de movimientos conservado"
            )

        messagebox.showinfo(
            "Material eliminado",
            mensaje,
            parent=root,
        )

    except Exception as error:

        traceback.print_exc()

        messagebox.showerror(
            "Error",
            f"No se pudo eliminar el material:\n\n"
            f"{error}",
            parent=root,
        )

    finally:

        importacion_en_curso = False


# ============================================================
# MOVIMIENTOS
# ============================================================

def actualizar_movimientos():

    if tabla_movimientos is None:
        return

    try:

        for x in tabla_movimientos.get_children():
            tabla_movimientos.delete(x)

    except Exception:
        return

    materiales = {
        m.get("id"): m
        for m in obtener_materiales_cache()
    }

    nombre_archivo = (
        inventario_seleccionado or ""
    )

    for mov in obtener_movimientos_cache():

        if (
            nombre_archivo
            and mov.get("archivo_origen")
            and mov.get("archivo_origen")
            != nombre_archivo
        ):
            continue

        m = materiales.get(
            mov.get("material_id")
        )

        try:

            tabla_movimientos.insert(
                "",
                "end",
                values=(
                    mov.get("fecha") or "",
                    m.get("codigo") if m else "",
                    m.get("material")
                    if m
                    else "Material eliminado",
                    mov.get("tipo") or "",
                    formatear_numero(
                        mov.get("cantidad", 0)
                    ),
                    formatear_numero(
                        mov.get("stock_anterior", 0)
                    ),
                    formatear_numero(
                        mov.get("stock_nuevo", 0)
                    ),
                    mov.get("usuario") or "",
                    mov.get("archivo_origen") or "",
                    mov.get("observaciones") or "",
                ),
            )

        except Exception:
            pass


# ============================================================
# REPORTES - UTILIDADES
# ============================================================

def obtener_movimientos_del_inventario():

    movimientos = obtener_movimientos_cache()

    nombre_archivo = (
        inventario_seleccionado or ""
    )

    resultado = []

    for movimiento in movimientos:

        if (
            nombre_archivo
            and movimiento.get("archivo_origen")
            and movimiento.get("archivo_origen")
            != nombre_archivo
        ):
            continue

        # Si existe documento_id, preferimos esa relación.
        documento_id = movimiento.get(
            "documento_id"
        )

        documento_actual_id = (
            obtener_documento_id_actual()
        )

        if (
            documento_actual_id is not None
            and documento_id is not None
        ):

            try:

                if int(documento_id) != int(
                    documento_actual_id
                ):
                    continue

            except Exception:
                pass

        resultado.append(movimiento)

    return resultado


def convertir_fecha_movimiento(valor):

    if not valor:
        return None

    texto = str(valor).strip()

    # ISO con fecha y hora
    try:

        return datetime.fromisoformat(
            texto.replace("Z", "+00:00")
        )

    except Exception:
        pass

    # ISO solamente fecha
    try:

        return datetime.strptime(
            texto[:10],
            "%Y-%m-%d"
        )

    except Exception:
        pass

    return None


def movimiento_dentro_de_fechas(
    movimiento,
    fecha_desde,
    fecha_hasta,
):

    fecha = convertir_fecha_movimiento(
        movimiento.get("fecha")
    )

    if fecha is None:
        return False

    # Quitamos timezone si solamente estamos
    # comparando fechas.
    fecha_solo = fecha.date()

    if fecha_desde:

        if fecha_solo < fecha_desde:
            return False

    if fecha_hasta:

        if fecha_solo > fecha_hasta:
            return False

    return True


# ============================================================
# REPORTE DE INVENTARIO
# ============================================================

def generar_reporte_inventario_excel():

    documento = obtener_documento_actual()

    if not documento:

        messagebox.showerror(
            "Reportes",
            "No se encontró el inventario seleccionado.",
            parent=root,
        )

        return

    try:

        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

    except ImportError:

        messagebox.showerror(
            "Reportes",
            "Falta instalar openpyxl.\n\n"
            "Ejecutá:\n\n"
            "pip install openpyxl",
            parent=root,
        )

        return

    materiales = obtener_materiales_seleccionados()

    nombre_documento = (
        documento.get("nombre")
        or "inventario"
    )

    nombre_base = Path(
        nombre_documento
    ).stem

    fecha_actual = datetime.now().strftime(
        "%Y-%m-%d_%H-%M"
    )

    ruta = filedialog.asksaveasfilename(
        parent=root,
        title="Guardar reporte de inventario",
        defaultextension=".xlsx",
        initialfile=(
            f"Reporte_Inventario_"
            f"{nombre_base}_"
            f"{fecha_actual}.xlsx"
        ),
        filetypes=[
            (
                "Excel",
                "*.xlsx",
            ),
            (
                "Todos los archivos",
                "*.*",
            ),
        ],
    )

    if not ruta:
        return

    try:

        wb = Workbook()

        ws = wb.active
        ws.title = "Inventario"

        # ----------------------------------------------------
        # CABECERA
        # ----------------------------------------------------

        ws["A1"] = "INVENTARIO MATERIAL NAVAL"

        ws["A1"].font = Font(
            bold=True,
            size=16,
        )

        ws["A2"] = "Inventario:"
        ws["B2"] = nombre_documento

        ws["A3"] = "Fecha del reporte:"
        ws["B3"] = datetime.now().strftime(
            "%d/%m/%Y %H:%M"
        )

        # ----------------------------------------------------
        # ENCABEZADOS
        # ----------------------------------------------------

        encabezados = [
            "Código",
            "Material",
            "Cantidad",
            "Unidad",
            "Categoría",
            "Ubicación",
            "Observaciones",
            "Estado",
        ]

        fila_encabezado = 5

        for columna, encabezado in enumerate(
            encabezados,
            start=1,
        ):

            celda = ws.cell(
                row=fila_encabezado,
                column=columna,
            )

            celda.value = encabezado

            celda.font = Font(
                bold=True
            )

            celda.alignment = Alignment(
                horizontal="center"
            )

        # ----------------------------------------------------
        # DATOS
        # ----------------------------------------------------

        fila = fila_encabezado + 1

        for material in materiales:

            try:

                cantidad = float(
                    material.get(
                        "cantidad",
                        0,
                    )
                    or 0
                )

            except Exception:

                cantidad = 0

            estado = (
                "HAY STOCK"
                if cantidad > 0
                else "SIN STOCK"
            )

            valores = [
                material.get("codigo") or "",
                material.get("material") or "",
                cantidad,
                material.get("unidad") or "",
                material.get("categoria") or "",
                material.get("ubicacion") or "",
                material.get("observaciones") or "",
                estado,
            ]

            for columna, valor in enumerate(
                valores,
                start=1,
            ):

                ws.cell(
                    row=fila,
                    column=columna,
                    value=valor,
                )

            fila += 1

        # ----------------------------------------------------
        # ANCHOS
        # ----------------------------------------------------

        anchos = {
            "A": 18,
            "B": 40,
            "C": 14,
            "D": 14,
            "E": 20,
            "F": 25,
            "G": 40,
            "H": 15,
        }

        for columna, ancho in anchos.items():

            ws.column_dimensions[
                columna
            ].width = ancho

        ws.freeze_panes = "A6"

        # ----------------------------------------------------
        # RESUMEN
        # ----------------------------------------------------

        fila_resumen = fila + 2

        total_materiales = len(
            materiales
        )

        con_stock = 0
        sin_stock = 0
        cantidad_total = 0

        for material in materiales:

            try:

                cantidad = float(
                    material.get(
                        "cantidad",
                        0,
                    )
                    or 0
                )

            except Exception:

                cantidad = 0

            cantidad_total += cantidad

            if cantidad > 0:
                con_stock += 1
            else:
                sin_stock += 1

        ws.cell(
            row=fila_resumen,
            column=1,
            value="Resumen",
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=fila_resumen + 1,
            column=1,
            value="Total materiales",
        )

        ws.cell(
            row=fila_resumen + 1,
            column=2,
            value=total_materiales,
        )

        ws.cell(
            row=fila_resumen + 2,
            column=1,
            value="Con stock",
        )

        ws.cell(
            row=fila_resumen + 2,
            column=2,
            value=con_stock,
        )

        ws.cell(
            row=fila_resumen + 3,
            column=1,
            value="Sin stock",
        )

        ws.cell(
            row=fila_resumen + 3,
            column=2,
            value=sin_stock,
        )

        ws.cell(
            row=fila_resumen + 4,
            column=1,
            value="Cantidad total",
        )

        ws.cell(
            row=fila_resumen + 4,
            column=2,
            value=cantidad_total,
        )

        wb.save(ruta)

        messagebox.showinfo(
            "Reporte generado",
            "El reporte de inventario se generó correctamente.\n\n"
            f"Archivo:\n{ruta}",
            parent=root,
        )

    except Exception as error:

        traceback.print_exc()

        messagebox.showerror(
            "Error",
            "No se pudo generar el reporte:\n\n"
            f"{error}",
            parent=root,
        )


# ============================================================
# REPORTE DE MOVIMIENTOS
# ============================================================

def generar_reporte_movimientos_excel(
    ventana=None,
    fecha_desde=None,
    fecha_hasta=None,
):

    documento = obtener_documento_actual()

    if not documento:

        messagebox.showerror(
            "Reportes",
            "No se encontró el inventario seleccionado.",
            parent=root,
        )

        return

    try:

        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

    except ImportError:

        messagebox.showerror(
            "Reportes",
            "Falta instalar openpyxl.\n\n"
            "Ejecutá:\n\n"
            "pip install openpyxl",
            parent=root,
        )

        return

    movimientos = (
        obtener_movimientos_del_inventario()
    )

    movimientos_filtrados = []

    for movimiento in movimientos:

        if movimiento_dentro_de_fechas(
            movimiento,
            fecha_desde,
            fecha_hasta,
        ):

            movimientos_filtrados.append(
                movimiento
            )

    nombre_documento = (
        documento.get("nombre")
        or "inventario"
    )

    nombre_base = Path(
        nombre_documento
    ).stem

    fecha_actual = datetime.now().strftime(
        "%Y-%m-%d_%H-%M"
    )

    ruta = filedialog.asksaveasfilename(
        parent=ventana or root,
        title="Guardar reporte de movimientos",
        defaultextension=".xlsx",
        initialfile=(
            f"Reporte_Movimientos_"
            f"{nombre_base}_"
            f"{fecha_actual}.xlsx"
        ),
        filetypes=[
            (
                "Excel",
                "*.xlsx",
            ),
            (
                "Todos los archivos",
                "*.*",
            ),
        ],
    )

    if not ruta:
        return

    try:

        wb = Workbook()

        ws = wb.active
        ws.title = "Movimientos"

        # ----------------------------------------------------
        # CABECERA
        # ----------------------------------------------------

        ws["A1"] = (
            "HISTORIAL DE MOVIMIENTOS"
        )

        ws["A1"].font = Font(
            bold=True,
            size=16,
        )

        ws["A2"] = "Inventario:"
        ws["B2"] = nombre_documento

        ws["A3"] = "Generado:"
        ws["B3"] = datetime.now().strftime(
            "%d/%m/%Y %H:%M"
        )

        ws["A4"] = "Desde:"
        ws["B4"] = (
            fecha_desde.strftime("%d/%m/%Y")
            if fecha_desde
            else "Todos"
        )

        ws["A5"] = "Hasta:"
        ws["B5"] = (
            fecha_hasta.strftime("%d/%m/%Y")
            if fecha_hasta
            else "Todos"
        )

        # ----------------------------------------------------
        # ENCABEZADOS
        # ----------------------------------------------------

        encabezados = [
            "Fecha",
            "Código",
            "Material",
            "Tipo",
            "Cantidad",
            "Stock anterior",
            "Stock nuevo",
            "Usuario",
            "Archivo",
            "Observaciones",
        ]

        fila_encabezado = 7

        for columna, encabezado in enumerate(
            encabezados,
            start=1,
        ):

            celda = ws.cell(
                row=fila_encabezado,
                column=columna,
            )

            celda.value = encabezado

            celda.font = Font(
                bold=True
            )

            celda.alignment = Alignment(
                horizontal="center"
            )

        # ----------------------------------------------------
        # MATERIALES
        # ----------------------------------------------------

        materiales_por_id = {
            m.get("id"): m
            for m in obtener_materiales_cache()
        }

        fila = fila_encabezado + 1

        total_entradas = 0
        total_salidas = 0

        for movimiento in movimientos_filtrados:

            material = materiales_por_id.get(
                movimiento.get(
                    "material_id"
                )
            )

            tipo = (
                movimiento.get("tipo")
                or ""
            )

            try:

                cantidad = float(
                    movimiento.get(
                        "cantidad",
                        0,
                    )
                    or 0
                )

            except Exception:

                cantidad = 0

            if tipo.upper() in (
                "ENTRADA",
                "INGRESO",
                "ALTA",
            ):

                total_entradas += cantidad

            elif tipo.upper() in (
                "SALIDA",
                "RETIRO",
               "BAJA",
            ):

                total_salidas += cantidad

            valores = [
                movimiento.get("fecha") or "",
                (
                    material.get("codigo")
                    if material
                    else ""
                ),
                (
                    material.get("material")
                    if material
                    else "Material eliminado"
                ),
                tipo,
                cantidad,
                movimiento.get(
                    "stock_anterior"
                ) or 0,
                movimiento.get(
                    "stock_nuevo"
                ) or 0,
                movimiento.get("usuario") or "",
                movimiento.get(
                    "archivo_origen"
                ) or "",
                movimiento.get(
                    "observaciones"
                ) or "",
            ]

            for columna, valor in enumerate(
                valores,
                start=1,
            ):

                ws.cell(
                    row=fila,
                    column=columna,
                    value=valor,
                )

            fila += 1

        # ----------------------------------------------------
        # ANCHOS
        # ----------------------------------------------------

        anchos = {
            "A": 25,
            "B": 18,
            "C": 40,
            "D": 15,
            "E": 14,
            "F": 18,
            "G": 18,
            "H": 15,
            "I": 35,
            "J": 40,
        }

        for columna, ancho in anchos.items():

            ws.column_dimensions[
                columna
            ].width = ancho

        ws.freeze_panes = "A8"

        # ----------------------------------------------------
        # RESUMEN
        # ----------------------------------------------------

        fila_resumen = fila + 2

        ws.cell(
            row=fila_resumen,
            column=1,
            value="Resumen",
        ).font = Font(
            bold=True
        )

        ws.cell(
            row=fila_resumen + 1,
            column=1,
            value="Movimientos registrados",
        )

        ws.cell(
            row=fila_resumen + 1,
            column=2,
            value=len(
                movimientos_filtrados
            ),
        )

        ws.cell(
            row=fila_resumen + 2,
            column=1,
            value="Total entradas",
        )

        ws.cell(
            row=fila_resumen + 2,
            column=2,
            value=total_entradas,
        )

        ws.cell(
            row=fila_resumen + 3,
            column=1,
            value="Total salidas",
        )

        ws.cell(
            row=fila_resumen + 3,
            column=2,
            value=total_salidas,
        )

        wb.save(ruta)

        if ventana is not None:

            try:
                ventana.destroy()
            except Exception:
                pass

        messagebox.showinfo(
            "Reporte generado",
            "El reporte de movimientos se generó correctamente.\n\n"
            f"Movimientos incluidos: "
            f"{len(movimientos_filtrados)}\n\n"
            f"Archivo:\n{ruta}",
            parent=root,
        )

    except Exception as error:

        traceback.print_exc()

        messagebox.showerror(
            "Error",
            "No se pudo generar el reporte:\n\n"
            f"{error}",
            parent=ventana or root,
        )


# ============================================================
# VENTANA DE REPORTES
# ============================================================

def abrir_reportes():

    if not inventario_seleccionado:

        messagebox.showwarning(
            "Reportes",
            "Primero seleccioná un inventario.",
            parent=root,
        )

        return

    ventana = tk.Toplevel(root)

    ventana.title(
        "Reportes"
    )

    ventana.geometry(
        "560x470"
    )

    ventana.transient(root)
    ventana.grab_set()

    marco = ttk.Frame(
        ventana,
        padding=25,
    )

    marco.pack(
        fill="both",
        expand=True,
    )

    ttk.Label(
        marco,
        text="📊 Reportes",
        font=(
            "Segoe UI",
            18,
            "bold",
        ),
    ).pack(
        pady=(5, 8)
    )

    ttk.Label(
        marco,
        text=(
            f"Inventario: "
            f"{inventario_seleccionado}"
        ),
        foreground=COLOR_AZUL,
        font=(
            "Segoe UI",
            10,
            "bold",
        ),
    ).pack(
        pady=(0, 20)
    )

    # ========================================================
    # REPORTE INVENTARIO
    # ========================================================

    ttk.Label(
        marco,
        text="Inventario actual",
        font=(
            "Segoe UI",
            11,
            "bold",
        ),
    ).pack(
        anchor="w"
    )

    ttk.Label(
        marco,
        text=(
            "Genera un Excel con todos los materiales "
            "del inventario seleccionado."
        ),
        foreground=COLOR_TEXTO_SECUNDARIO,
    ).pack(
        anchor="w",
        pady=(2, 7),
    )

    ttk.Button(
        marco,
        text="📦 Generar reporte de inventario",
        command=generar_reporte_inventario_excel,
    ).pack(
        fill="x",
        pady=(0, 18),
        ipady=4,
    )

    # ========================================================
    # REPORTE MOVIMIENTOS
    # ========================================================

    ttk.Label(
        marco,
        text="Movimientos",
        font=(
            "Segoe UI",
            11,
            "bold",
        ),
    ).pack(
        anchor="w"
    )

    ttk.Label(
        marco,
        text=(
            "Podés generar el historial completo "
            "o seleccionar un período."
        ),
        foreground=COLOR_TEXTO_SECUNDARIO,
    ).pack(
        anchor="w",
        pady=(2, 8),
    )

    frame_fechas = ttk.Frame(
        marco
    )

    frame_fechas.pack(
        fill="x",
        pady=(0, 10),
    )

    ttk.Label(
        frame_fechas,
        text="Desde:",
    ).grid(
        row=0,
        column=0,
        sticky="w",
        padx=(0, 5),
    )

    entrada_desde = ttk.Entry(
        frame_fechas,
        width=15,
    )

    entrada_desde.grid(
        row=0,
        column=1,
        sticky="w",
        padx=(0, 15),
    )

    ttk.Label(
        frame_fechas,
        text="Hasta:",
    ).grid(
        row=0,
        column=2,
        sticky="w",
        padx=(0, 5),
    )

    entrada_hasta = ttk.Entry(
        frame_fechas,
        width=15,
    )

    entrada_hasta.grid(
        row=0,
        column=3,
        sticky="w",
    )

    ttk.Label(
        marco,
        text="Formato: DD/MM/AAAA",
        foreground=COLOR_TEXTO_SECUNDARIO,
    ).pack(
        anchor="w",
        pady=(0, 8),
    )

    def generar_movimientos():

        texto_desde = (
            entrada_desde
            .get()
            .strip()
        )

        texto_hasta = (
            entrada_hasta
            .get()
            .strip()
        )

        fecha_desde = None
        fecha_hasta = None

        if texto_desde:

            try:

                fecha_desde = datetime.strptime(
                    texto_desde,
                    "%d/%m/%Y",
                ).date()

            except ValueError:

                messagebox.showerror(
                    "Fecha",
                    "La fecha 'Desde' no es válida.\n\n"
                    "Usá el formato DD/MM/AAAA.",
                    parent=ventana,
                )

                return

        if texto_hasta:

            try:

                fecha_hasta = datetime.strptime(
                    texto_hasta,
                    "%d/%m/%Y",
                ).date()

            except ValueError:

                messagebox.showerror(
                    "Fecha",
                    "La fecha 'Hasta' no es válida.\n\n"
                    "Usá el formato DD/MM/AAAA.",
                    parent=ventana,
                )

                return

        if (
            fecha_desde
            and fecha_hasta
            and fecha_desde > fecha_hasta
        ):

            messagebox.showerror(
                "Fechas",
                "La fecha 'Desde' no puede ser posterior "
                "a la fecha 'Hasta'.",
                parent=ventana,
            )

            return

        generar_reporte_movimientos_excel(
            ventana=ventana,
            fecha_desde=fecha_desde,
            fecha_hasta=fecha_hasta,
        )

    ttk.Button(
        marco,
        text="📜 Generar reporte de movimientos",
        command=generar_movimientos,
    ).pack(
        fill="x",
        pady=(0, 18),
        ipady=4,
    )

    ttk.Button(
        marco,
        text="Cerrar",
        command=ventana.destroy,
    ).pack(
        pady=(0, 5)
    )


# ============================================================
# ACTUALIZAR TODO
# ============================================================

def actualizar_todo(forzar=False):

    global actualizacion_en_curso
    global firma_datos_sincronizados

    with lock_actualizacion:

        if actualizacion_en_curso:
            return

        actualizacion_en_curso = True

    def trabajo():

        global actualizacion_en_curso
        global firma_datos_sincronizados

        try:

            if lbl_progreso:

                root.after(
                    0,
                    lambda: lbl_progreso.config(
                        text="⏳ Actualizando..."
                    ),
                )

            (
                materiales,
                documentos,
                movimientos,
            ) = cargar_datos_supabase()

            firma_datos_sincronizados = (
                generar_firma_sincronizacion(
                    materiales,
                    documentos,
                    movimientos,
                )
            )

            def refrescar():

                global actualizacion_en_curso

                try:

                    refrescar_pantalla_seleccion()

                    actualizar_tabla()
                    actualizar_movimientos()

                    if lbl_estado:

                        lbl_estado.config(
                            text="🟢 Sincronizado"
                        )

                    if lbl_progreso:

                        lbl_progreso.config(
                            text="✓ Actualizado"
                        )

                finally:

                    with lock_actualizacion:
                        actualizacion_en_curso = False

            root.after(
                0,
                refrescar,
            )

        except Exception:

            traceback.print_exc()

            def fallo():

                global actualizacion_en_curso

                with lock_actualizacion:
                    actualizacion_en_curso = False

                if lbl_estado:

                    lbl_estado.config(
                        text="🔴 Sin conexión"
                    )

                if lbl_progreso:

                    lbl_progreso.config(
                        text="Error al actualizar"
                    )

            root.after(
                0,
                fallo,
            )

    threading.Thread(
        target=trabajo,
        daemon=True,
    ).start()


# ============================================================
# SINCRONIZACIÓN AUTOMÁTICA
# ============================================================

def sincronizar_automaticamente():

    global sincronizacion_en_curso
    global firma_datos_sincronizados
    global cache_materiales
    global cache_documentos
    global cache_movimientos

    if (
        root is None
        or importacion_en_curso
        or pausar_sincronizacion
    ):
        return

    with lock_sincronizacion:

        if sincronizacion_en_curso:
            return

        sincronizacion_en_curso = True

    def trabajo():

        global sincronizacion_en_curso
        global firma_datos_sincronizados
        global cache_materiales
        global cache_documentos
        global cache_movimientos

        try:

            with lock_actualizacion:

                if actualizacion_en_curso:
                    return

            materiales = (
                obtener_materiales()
                or []
            )

            documentos = (
                obtener_documentos()
                or []
            )

            movimientos = (
                obtener_movimientos(100)
                or []
            )

            firma = (
                generar_firma_sincronizacion(
                    materiales,
                    documentos,
                    movimientos,
                )
            )

            cambio = (
                firma_datos_sincronizados is None
                or firma
                != firma_datos_sincronizados
            )

            firma_datos_sincronizados = firma

            if cambio:

                with cache_lock:

                    cache_materiales = materiales
                    cache_documentos = documentos
                    cache_movimientos = movimientos

                root.after(
                    0,
                    actualizar_interfaz_por_sincronizacion,
                )

            else:

                root.after(
                    0,
                    lambda: actualizar_estado_sync(
                        True,
                        False,
                    ),
                )

        except Exception as error:

            print(
                "❌ Error en sincronización automática:",
                error,
            )

            traceback.print_exc()

            root.after(
                0,
                lambda: actualizar_estado_sync(
                    False,
                    False,
                ),
            )

        finally:

            with lock_sincronizacion:
                sincronizacion_en_curso = False

    threading.Thread(
        target=trabajo,
        daemon=True,
    ).start()


def actualizar_interfaz_por_sincronizacion():

    try:

        refrescar_pantalla_seleccion()

        actualizar_tabla()

        actualizar_movimientos()

        if lbl_estado:

            lbl_estado.config(
                text="🟢 Sincronizado"
            )

        if lbl_progreso:

            lbl_progreso.config(
                text="🔄 Cambio detectado — sincronizado"
            )

    except Exception:

        traceback.print_exc()


def actualizar_estado_sync(
    conectado,
    hubo_cambio,
):

    if lbl_estado:

        lbl_estado.config(
            text=(
                "🟢 Sincronizado"
                if conectado
                else "🔴 Sin conexión"
            )
        )

    if (
        hubo_cambio
        and lbl_progreso
    ):

        lbl_progreso.config(
            text="🔄 Datos sincronizados"
        )


def monitor_sincronizacion():

    print(
        "🔄 Monitor de sincronización iniciado."
    )

    while True:

        try:
            sincronizar_automaticamente()

        except Exception:
            traceback.print_exc()

        time.sleep(
            INTERVALO_SINCRONIZACION
            / 1000
        )


# ============================================================
# IMPORTACIÓN WORD
# ============================================================

def ejecutar_importacion_word(
    reescaneo_completo=False,
):

    global importacion_en_curso
    global pausar_sincronizacion
    global firma_datos_sincronizados
    global cache_materiales
    global cache_documentos
    global cache_movimientos

    with lock_importacion:

        if importacion_en_curso:
            return

        importacion_en_curso = True

    pausar_sincronizacion = True

    if lbl_progreso:

        lbl_progreso.config(
            text="⏳ Importando documentos Word..."
        )

    if lbl_estado:

        lbl_estado.config(
            text="🟡 Importando Word..."
        )

    def trabajo():

        global importacion_en_curso
        global pausar_sincronizacion
        global firma_datos_sincronizados
        global cache_materiales
        global cache_documentos
        global cache_movimientos

        resultado = None

        try:

            resultado = (
                importar_word.importar_todos(
                    reescaneo_completo=
                    reescaneo_completo
                )
            )

            invalidar_cache()

            materiales = (
                obtener_materiales()
                or []
            )

            documentos = (
                obtener_documentos()
                or []
            )

            movimientos = (
                obtener_movimientos(100)
                or []
            )

            with cache_lock:

                cache_materiales = materiales
                cache_documentos = documentos
                cache_movimientos = movimientos

            firma_datos_sincronizados = (
                generar_firma_sincronizacion(
                    materiales,
                    documentos,
                    movimientos,
                )
            )

            print(
                "RESULTADO:",
                resultado,
            )

        except Exception as error:

            traceback.print_exc()

            root.after(
                0,
                lambda error=error:
                messagebox.showerror(
                    "Error",
                    "Ocurrió un error durante "
                    "la importación:\n\n"
                    f"{error}",
                    parent=root,
                ),
            )

        finally:

            pausar_sincronizacion = False
            importacion_en_curso = False

            root.after(
                0,
                actualizar_interfaz_por_sincronizacion,
            )

            if resultado is not None:

                root.after(
                    150,
                    lambda:
                    mostrar_resultado_importacion(
                        resultado
                    ),
                )

    threading.Thread(
        target=trabajo,
        daemon=True,
    ).start()


def construir_mensaje_importacion(
    resultado
):

    if not isinstance(
        resultado,
        dict,
    ):
        return str(resultado)

    campos = [
        ("nuevos", "Nuevos"),
        ("modificados", "Modificados"),
        ("actualizados", "Actualizados"),
        ("sin_cambios", "Sin cambios"),
        ("eliminados", "Eliminados"),
        ("vacios", "Vacíos"),
        ("filas", "Filas procesadas"),
        ("items", "Items"),
        ("errores", "Errores"),
    ]

    return "\n".join(
        f"{nombre}: "
        f"{resultado.get(clave, 0)}"
        for clave, nombre in campos
        if clave in resultado
    ) or str(resultado)


def mostrar_resultado_importacion(
    resultado
):

    if lbl_progreso:

        lbl_progreso.config(
            text="✓ Word sincronizado"
        )

    messagebox.showinfo(
        "Importación finalizada",
        construir_mensaje_importacion(
            resultado
        ),
        parent=root,
    )


def importar_word_manual():

    if importacion_en_curso:

        messagebox.showinfo(
            "Importación",
            "Ya hay una importación en curso.",
            parent=root,
        )

        return

    ejecutar_importacion_word(False)


def reescaneo_completo():

    if importacion_en_curso:

        messagebox.showinfo(
            "Importación",
            "Ya hay una importación en curso.",
            parent=root,
        )

        return

    if not messagebox.askyesno(
        "Reescaneo completo",
        "Se van a revisar nuevamente "
        "todos los archivos Word.\n\n"
        "Esto puede tardar unos minutos.\n\n"
        "¿Continuar?",
        parent=root,
    ):
        return

    ejecutar_importacion_word(True)


# ============================================================
# MONITOR DE WORD
# ============================================================

def obtener_estado_archivos_word():

    estado = {}

    CARPETA_DOCUMENTOS.mkdir(
        exist_ok=True
    )

    for ruta in CARPETA_DOCUMENTOS.glob(
        "*.docx"
    ):

        if ruta.name.startswith("~$"):
            continue

        try:

            estado[
                str(ruta.resolve())
            ] = ruta.stat().st_mtime

        except Exception:
            pass

    return estado


def monitor_word():

    global estado_archivos_word

    print(
        "📄 Monitor de Word iniciado."
    )

    while True:

        try:

            nuevo = (
                obtener_estado_archivos_word()
            )

            if nuevo != estado_archivos_word:

                estado_archivos_word = nuevo

                if not importacion_en_curso:

                    ejecutar_importacion_word(
                        False
                    )

        except Exception:
            traceback.print_exc()

        time.sleep(
            INTERVALO_MONITOR
            / 1000
        )


# ============================================================
# ESTILO
# ============================================================

def aplicar_estilo():

    estilo = ttk.Style()

    try:
        estilo.theme_use("clam")

    except Exception:
        pass

    estilo.configure(
        ".",
        font=("Segoe UI", 10),
    )

    estilo.configure(
        "TFrame",
        background=COLOR_FONDO,
    )

    estilo.configure(
        "TLabel",
        background=COLOR_FONDO,
        foreground=COLOR_TEXTO,
    )

    estilo.configure(
        "Title.TLabel",
        background=COLOR_AZUL_OSCURO,
        foreground="white",
        font=("Segoe UI", 18, "bold"),
    )

    estilo.configure(
        "Subtitle.TLabel",
        background=COLOR_AZUL_OSCURO,
        foreground="#dce8f1",
        font=("Segoe UI", 9),
    )

    estilo.configure(
        "TButton",
        padding=(12, 7),
        font=("Segoe UI", 9, "bold"),
    )

    estilo.configure(
        "Treeview",
        background="white",
        foreground=COLOR_TEXTO,
        fieldbackground="white",
        rowheight=30,
        font=("Segoe UI", 10),
        borderwidth=0,
    )

    estilo.configure(
        "Treeview.Heading",
        background=COLOR_AZUL_OSCURO,
        foreground="white",
        font=("Segoe UI", 10, "bold"),
        padding=7,
    )

    estilo.map(
        "Treeview",
        background=[
            ("selected", COLOR_AZUL)
        ],
        foreground=[
            ("selected", "white")
        ],
    )


# ============================================================
# CABECERA
# ============================================================

def crear_cabecera():

    cab = tk.Frame(
        root,
        bg=COLOR_AZUL_OSCURO,
        height=75,
    )

    cab.pack(fill="x")
    cab.pack_propagate(False)

    tit = tk.Frame(
        cab,
        bg=COLOR_AZUL_OSCURO,
    )

    tit.pack(
        side="left",
        padx=20,
    )

    ttk.Label(
        tit,
        text="⚓ INVENTARIO MATERIAL NAVAL",
        style="Title.TLabel",
    ).pack(anchor="w")

    ttk.Label(
        tit,
        text="Sistema de gestión y control de material",
        style="Subtitle.TLabel",
    ).pack(anchor="w")

    est = tk.Frame(
        cab,
        bg=COLOR_AZUL_OSCURO,
    )

    est.pack(
        side="right",
        padx=20,
    )

    global lbl_estado

    lbl_estado = tk.Label(
        est,
        text="🟡 Conectando...",
        bg=COLOR_AZUL_OSCURO,
        fg="white",
        font=("Segoe UI", 10, "bold"),
    )

    lbl_estado.pack(
        side="right"
    )


# ============================================================
# SELECCIÓN DEL INVENTARIO
# ============================================================

def seleccionar_inventario(documento):

    global inventario_seleccionado
    global inventario_seleccionado_id

    if isinstance(documento, str):

        nombre = documento

        encontrado = None

        for d in obtener_documentos_cache():

            if d.get("nombre") == nombre:

                encontrado = d
                break

        if encontrado is None:

            messagebox.showerror(
                "Inventario",
                "No se pudo encontrar el documento seleccionado.",
                parent=root,
            )

            return

        documento = encontrado

    if not isinstance(documento, dict):

        messagebox.showerror(
            "Inventario",
            "El documento seleccionado no es válido.",
            parent=root,
        )

        return

    documento_id = documento.get("id")
    nombre = documento.get("nombre")

    if documento_id is None:

        messagebox.showerror(
            "Inventario",
            "El documento seleccionado no tiene ID.",
            parent=root,
        )

        return

    if not nombre:

        nombre = (
            Path(
                str(
                    documento.get(
                        "ruta",
                        "",
                    )
                )
            ).name
        )

    inventario_seleccionado_id = documento_id
    inventario_seleccionado = nombre

    print(
        "📂 Inventario seleccionado:"
    )

    print(
        "   ID:",
        inventario_seleccionado_id,
    )

    print(
        "   Nombre:",
        inventario_seleccionado,
    )

    print(
        "   Ruta:",
        documento.get("ruta"),
    )

    if entrada_busqueda:

        try:
            entrada_busqueda.delete(
                0,
                tk.END,
            )
        except Exception:
            pass

    construir_pantalla_inventario()


def mostrar_estado_selector(mensaje, detalle=None, error=False):

    if marco_selector is None:
        return

    for widget in marco_selector.winfo_children():
        try:
            widget.destroy()
        except Exception:
            pass

    cont = tk.Frame(
        marco_selector,
        bg=COLOR_FONDO,
    )

    cont.pack(
        fill="both",
        expand=True,
    )

    simbolo = "🔴" if error else "⏳"

    tk.Label(
        cont,
        text=f"{simbolo} {mensaje}",
        bg=COLOR_FONDO,
        fg=COLOR_TEXTO,
        font=("Segoe UI", 16, "bold"),
    ).pack(pady=(120, 8))

    if detalle:
        tk.Label(
            cont,
            text=detalle,
            bg=COLOR_FONDO,
            fg=COLOR_TEXTO_SECUNDARIO,
            font=("Segoe UI", 10),
        ).pack()


# ============================================================
# REFRESCAR SELECTOR
# ============================================================

def refrescar_pantalla_seleccion():

    if inventario_seleccionado:
        return

    if (
        marco_selector is None
        or not marco_selector.winfo_exists()
    ):
        return

    if not datos_iniciales_cargados:
        if error_carga_inicial:
            mostrar_estado_selector(
                "No se pudieron cargar los inventarios",
                str(error_carga_inicial),
                error=True,
            )
        else:
            mostrar_estado_selector(
                "Cargando inventarios...",
                "Conectando con la base de datos y cargando los documentos.",
            )
        return

    for widget in marco_selector.winfo_children():
        try:
            widget.destroy()
        except Exception:
            pass

    construir_opciones_documentos(
        marco_selector
    )


# ============================================================
# TARJETAS DE DOCUMENTOS
# ============================================================

def construir_opciones_documentos(parent):

    documentos = [
        d
        for d in obtener_documentos_cache()
        if d.get("nombre")
    ]

    ttk.Label(
        parent,
        text="Seleccioná el inventario",
        font=("Segoe UI", 15, "bold"),
    ).pack(
        pady=(20, 5)
    )

    ttk.Label(
        parent,
        text="Cada archivo Word administra su propio stock.",
        foreground=COLOR_TEXTO_SECUNDARIO,
    ).pack(
        pady=(0, 18)
    )

    cont = tk.Frame(
        parent,
        bg=COLOR_FONDO,
    )

    cont.pack(
        fill="both",
        expand=True,
        padx=40,
        pady=10,
    )

    if not documentos:

        ttk.Label(
            cont,
            text="No se encontraron archivos Word en la carpeta documentos.",
            font=("Segoe UI", 11),
        ).pack(
            pady=40
        )

        return

    columnas = 2

    for i, doc in enumerate(documentos):

        nombre = doc.get(
            "nombre"
        )

        tarjeta = tk.Frame(
            cont,
            bg=COLOR_PANEL,
            highlightbackground=COLOR_BORDE,
            highlightthickness=1,
            padx=18,
            pady=15,
            cursor="hand2",
        )

        tarjeta.grid(
            row=i // columnas,
            column=i % columnas,
            sticky="nsew",
            padx=10,
            pady=10,
        )

        cont.grid_columnconfigure(
            i % columnas,
            weight=1,
        )

        cont.grid_rowconfigure(
            i // columnas,
            weight=1,
        )

        def abrir_documento(
            evento=None,
            documento=doc,
        ):
            seleccionar_inventario(
                documento
            )

        icono = tk.Label(
            tarjeta,
            text="📄",
            bg=COLOR_PANEL,
            fg=COLOR_AZUL,
            font=("Segoe UI", 24),
            cursor="hand2",
        )

        icono.pack(
            side="left",
            padx=(0, 15),
        )

        info = tk.Frame(
            tarjeta,
            bg=COLOR_PANEL,
        )

        info.pack(
            side="left",
            fill="x",
            expand=True,
        )

        etiqueta_nombre = tk.Label(
            info,
            text=nombre,
            bg=COLOR_PANEL,
            fg=COLOR_TEXTO,
            font=("Segoe UI", 11, "bold"),
            anchor="w",
            cursor="hand2",
        )

        etiqueta_nombre.pack(
            fill="x"
        )

        etiqueta_descripcion = tk.Label(
            info,
            text="Abrir inventario",
            bg=COLOR_PANEL,
            fg=COLOR_TEXTO_SECUNDARIO,
            font=("Segoe UI", 9),
            anchor="w",
            cursor="hand2",
        )

        etiqueta_descripcion.pack(
            fill="x",
            pady=(2, 8),
        )

        boton = ttk.Button(
            info,
            text="Abrir",
            command=abrir_documento,
        )

        boton.pack(
            anchor="w"
        )

        tarjeta.bind(
            "<Button-1>",
            abrir_documento,
        )

        icono.bind(
            "<Button-1>",
            abrir_documento,
        )

        etiqueta_nombre.bind(
            "<Button-1>",
            abrir_documento,
        )

        etiqueta_descripcion.bind(
            "<Button-1>",
            abrir_documento,
        )


# ============================================================
# MOSTRAR SELECTOR
# ============================================================

def mostrar_pantalla_seleccion():

    global marco_selector
    global marco_contenido

    global tabla
    global tabla_movimientos
    global entrada_busqueda

    if marco_contenido is not None:

        try:
            marco_contenido.destroy()
        except Exception:
            pass

        marco_contenido = None

    tabla = None
    tabla_movimientos = None
    entrada_busqueda = None

    if marco_selector is not None:

        try:
            marco_selector.destroy()
        except Exception:
            pass

    marco_selector = tk.Frame(
        root,
        bg=COLOR_FONDO,
    )

    marco_selector.pack(
        fill="both",
        expand=True,
    )

    refrescar_pantalla_seleccion()


# ============================================================
# VOLVER AL SELECTOR
# ============================================================

def volver_a_seleccion():

    global inventario_seleccionado
    global inventario_seleccionado_id

    inventario_seleccionado = None
    inventario_seleccionado_id = None

    mostrar_pantalla_seleccion()


# ============================================================
# CARGAR INVENTARIO SELECCIONADO
# ============================================================

def cargar_inventario_seleccionado():

    if not inventario_seleccionado:
        return

    documento = obtener_documento_actual()

    if documento is None:

        messagebox.showerror(
            "Inventario",
            "No se pudo encontrar el inventario seleccionado en la base de datos.",
            parent=root,
        )

        volver_a_seleccion()

        return

    try:

        actualizar_tabla()
        actualizar_movimientos()

        if lbl_progreso:

            lbl_progreso.config(
                text="✓ Inventario cargado"
            )

    except Exception as error:

        traceback.print_exc()

        if lbl_progreso:

            lbl_progreso.config(
                text="Error al cargar inventario"
            )

        messagebox.showerror(
            "Error",
            f"No se pudo cargar el inventario:\n\n{error}",
            parent=root,
        )


# ============================================================
# RELACIÓN DE TRÁNSITO
# ============================================================

def generar_relacion_transito_ui():

    material = obtener_material_seleccionado()

    if not material:
        return

    tipo_movimiento = simpledialog.askstring(
        "Relación de Tránsito",
        "Tipo de movimiento:",
        parent=root,
    )

    if tipo_movimiento is None:
        return

    destino = simpledialog.askstring(
        "Relación de Tránsito",
        "Destino:",
        parent=root,
    )

    if destino is None:
        return

    transporte = simpledialog.askstring(
        "Relación de Tránsito",
        "Transporte:",
        parent=root,
    )

    if transporte is None:
        return

    nombre = limpiar_texto(material.get("material")) or "material"
    nombre_archivo = f"RELACION DE TRANSITO_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    salida = filedialog.asksaveasfilename(
        parent=root,
        title="Guardar Relación de Tránsito",
        defaultextension=".xlsx",
        initialfile=nombre_archivo,
        filetypes=[("Excel", "*.xlsx")],
    )

    if not salida:
        return

    try:
        generar_relacion_transito(
            materiales=[material],
            salida=salida,
            tipo_movimiento=tipo_movimiento,
            destino=destino,
            transporte=transporte,
        )

        messagebox.showinfo(
            "Relación de Tránsito",
            f"La relación de tránsito de \"{nombre}\" se generó correctamente.\n\n"
            f"Archivo:\n{salida}",
            parent=root,
        )

    except RelacionTransitoError as error:
        traceback.print_exc()
        messagebox.showerror(
            "Relación de Tránsito",
            f"No se pudo generar la relación de tránsito:\n\n{error}",
            parent=root,
        )

    except Exception as error:
        traceback.print_exc()
        messagebox.showerror(
            "Relación de Tránsito",
            f"Ocurrió un error inesperado al generar la relación:\n\n{error}",
            parent=root,
        )


# ============================================================
# CONSTRUIR PANTALLA DEL INVENTARIO
# ============================================================

def construir_pantalla_inventario():

    global marco_contenido
    global marco_selector

    global tabla
    global tabla_movimientos
    global entrada_busqueda

    global lbl_total_materiales
    global lbl_con_stock
    global lbl_sin_stock
    global lbl_cantidad_total
    global lbl_progreso

    if marco_selector is not None:

        try:
            marco_selector.destroy()
        except Exception:
            pass

        marco_selector = None

    if marco_contenido is not None:

        try:
            marco_contenido.destroy()
        except Exception:
            pass

    tabla = None
    tabla_movimientos = None
    entrada_busqueda = None

    marco_contenido = tk.Frame(
        root,
        bg=COLOR_FONDO,
    )

    marco_contenido.pack(
        fill="both",
        expand=True,
    )

    fs = ttk.Frame(
        marco_contenido,
        padding=(15, 10),
    )

    fs.pack(
        fill="x"
    )

    ttk.Label(
        fs,
        text="Inventario:",
        font=("Segoe UI", 10, "bold"),
    ).pack(
        side="left",
        padx=(0, 8),
    )

    ttk.Label(
        fs,
        text=inventario_seleccionado or "",
        font=("Segoe UI", 10, "bold"),
        foreground=COLOR_AZUL,
    ).pack(
        side="left"
    )

    ttk.Button(
        fs,
        text="📂 Cambiar inventario",
        command=volver_a_seleccion,
    ).pack(
        side="right"
    )

    stats = tk.Frame(
        marco_contenido,
        bg=COLOR_AZUL_CLARO,
        highlightbackground=COLOR_BORDE,
        highlightthickness=1,
    )

    stats.pack(
        fill="x",
        padx=15,
        pady=(0, 7),
    )

    def stat(t):

        return tk.Label(
            stats,
            text=t,
            bg=COLOR_AZUL_CLARO,
            fg=COLOR_TEXTO,
            font=("Segoe UI", 10, "bold"),
            padx=15,
            pady=8,
        )

    lbl_total_materiales = stat(
        "Materiales: 0"
    )

    lbl_total_materiales.pack(
        side="left"
    )

    lbl_con_stock = stat(
        "Con stock: 0"
    )

    lbl_con_stock.pack(
        side="left"
    )

    lbl_sin_stock = stat(
        "Sin stock: 0"
    )

    lbl_sin_stock.pack(
        side="left"
    )

    lbl_cantidad_total = stat(
        "Cantidad total: 0"
    )

    lbl_cantidad_total.pack(
        side="left"
    )

    fb = ttk.Frame(
        marco_contenido,
        padding=(15, 5),
    )

    fb.pack(
        fill="x"
    )

    ttk.Label(
        fb,
        text="Buscar:",
        font=("Segoe UI", 10, "bold"),
    ).pack(
        side="left",
        padx=(0, 7),
    )

    entrada_busqueda = ttk.Entry(
        fb
    )

    entrada_busqueda.pack(
        side="left",
        fill="x",
        expand=True,
    )

    entrada_busqueda.bind(
        "<KeyRelease>",
        buscar,
    )

    # ========================================================
    # BOTONES
    # ========================================================

    buttons = ttk.Frame(
        marco_contenido,
        padding=(15, 5),
    )

    buttons.pack(
        fill="x"
    )

    for text, cmd in (
        ("➕ Nuevo", nuevo_material),
        ("✏️ Editar", editar_material),
        ("🗑️ Eliminar", eliminar_material),
        ("📥 Agregar", agregar_stock),
        ("📤 Retirar", retirar_stock),
        ("📊 Reportes", abrir_reportes),
        ("📋 Relación de Tránsito", generar_relacion_transito_ui),
        (
            "🔄 Actualizar",
            lambda: actualizar_todo(True),
        ),
    ):

        ttk.Button(
            buttons,
            text=text,
            command=cmd,
        ).pack(
            side="left",
            padx=3,
        )

    ttk.Button(
        buttons,
        text="📄 Importar Word",
        command=importar_word_manual,
    ).pack(
        side="right",
        padx=3,
    )

    ttk.Button(
        buttons,
        text="🧹 Reescaneo completo",
        command=reescaneo_completo,
    ).pack(
        side="right",
        padx=3,
    )

    lbl_progreso = ttk.Label(
        marco_contenido,
        text="Cargando inventario...",
        foreground=COLOR_TEXTO_SECUNDARIO,
    )

    lbl_progreso.pack(
        anchor="w",
        padx=18,
        pady=(2, 2),
    )

    ft = ttk.Frame(
        marco_contenido,
        padding=(15, 3),
    )

    ft.pack(
        fill="both",
        expand=True,
    )

    cols = (
        "codigo",
        "material",
        "cantidad",
        "unidad",
        "categoria",
        "ubicacion",
        "observaciones",
        "estado",
    )

    tabla = ttk.Treeview(
        ft,
        columns=cols,
        show="headings",
        selectmode="browse",
    )

    heads = {
        "codigo": "Código",
        "material": "Material",
        "cantidad": "Cantidad",
        "unidad": "Unidad",
        "categoria": "Categoría",
        "ubicacion": "Ubicación",
        "observaciones": "Observaciones",
        "estado": "Estado",
    }

    widths = {
        "codigo": 150,
        "material": 300,
        "cantidad": 100,
        "unidad": 100,
        "categoria": 150,
        "ubicacion": 180,
        "observaciones": 300,
        "estado": 150,
    }

    for c in cols:

        tabla.heading(
            c,
            text=heads[c],
        )

        tabla.column(
            c,
            width=widths[c],
            minwidth=70,
        )

    sv = ttk.Scrollbar(
        ft,
        orient="vertical",
        command=tabla.yview,
    )

    sh = ttk.Scrollbar(
        ft,
        orient="horizontal",
        command=tabla.xview,
    )

    tabla.configure(
        yscrollcommand=sv.set,
        xscrollcommand=sh.set,
    )

    tabla.grid(
        row=0,
        column=0,
        sticky="nsew",
    )

    sv.grid(
        row=0,
        column=1,
        sticky="ns",
    )

    sh.grid(
        row=1,
        column=0,
        sticky="ew",
    )

    ft.rowconfigure(
        0,
        weight=1,
    )

    ft.columnconfigure(
        0,
        weight=1,
    )

    tabla.tag_configure(
        "stock",
        background=COLOR_STOCK_FONDO,
    )

    tabla.tag_configure(
        "sin_stock",
        background=COLOR_SIN_STOCK_FONDO,
    )

    ttk.Label(
        marco_contenido,
        text="Últimos movimientos",
        font=("Segoe UI", 12, "bold"),
    ).pack(
        anchor="w",
        padx=15,
        pady=(8, 3),
    )

    fm = ttk.Frame(
        marco_contenido,
        height=180,
    )

    fm.pack(
        fill="x",
        padx=15,
        pady=(0, 10),
    )

    fm.pack_propagate(False)

    mc = (
        "fecha",
        "codigo",
        "material",
        "tipo",
        "cantidad",
        "anterior",
        "nuevo",
        "usuario",
        "archivo",
        "observaciones",
    )

    tabla_movimientos = ttk.Treeview(
        fm,
        columns=mc,
        show="headings",
    )

    mh = {
        "fecha": "Fecha",
        "codigo": "Código",
        "material": "Material",
        "tipo": "Tipo",
        "cantidad": "Cantidad",
        "anterior": "Stock anterior",
        "nuevo": "Stock nuevo",
        "usuario": "Usuario",
        "archivo": "Archivo",
        "observaciones": "Observaciones",
    }

    mw = {
        "fecha": 160,
        "codigo": 130,
        "material": 250,
        "tipo": 100,
        "cantidad": 100,
        "anterior": 110,
        "nuevo": 110,
        "usuario": 100,
        "archivo": 250,
        "observaciones": 250,
    }

    for c in mc:

        tabla_movimientos.heading(
            c,
            text=mh[c],
        )

        tabla_movimientos.column(
            c,
            width=mw[c],
            minwidth=80,
        )

    smv = ttk.Scrollbar(
        fm,
        orient="vertical",
        command=tabla_movimientos.yview,
    )

    smh = ttk.Scrollbar(
        fm,
        orient="horizontal",
        command=tabla_movimientos.xview,
    )

    tabla_movimientos.configure(
        yscrollcommand=smv.set,
        xscrollcommand=smh.set,
    )

    tabla_movimientos.grid(
        row=0,
        column=0,
        sticky="nsew",
    )

    smv.grid(
        row=0,
        column=1,
        sticky="ns",
    )

    smh.grid(
        row=1,
        column=0,
        sticky="ew",
    )

    fm.rowconfigure(
        0,
        weight=1,
    )

    fm.columnconfigure(
        0,
        weight=1,
    )

    root.after(
        100,
        cargar_inventario_seleccionado,
    )


# ============================================================
# CREAR INTERFAZ
# ============================================================

def crear_interfaz():

    global root
    global estado_archivos_word

    root = tk.Tk()

    root.title(
        NOMBRE_APP
    )

    root.geometry(
        "1500x900"
    )

    root.minsize(
        1100,
        700,
    )

    root.configure(
        bg=COLOR_FONDO
    )

    aplicar_estilo()

    crear_cabecera()

    mostrar_pantalla_seleccion()

    def iniciar():

        global datos_iniciales_cargados
        global error_carga_inicial

        try:

            conectado = probar_conexion()

            root.after(
                0,
                lambda:
                lbl_estado.config(
                    text=(
                        "🟢 Conectado"
                        if conectado
                        else "🔴 Sin conexión"
                    )
                ),
            )

            cargar_datos_supabase()

            datos_iniciales_cargados = True
            error_carga_inicial = None

            root.after(
                0,
                refrescar_pantalla_seleccion,
            )

        except Exception as error:

            global datos_iniciales_cargados
            global error_carga_inicial

            datos_iniciales_cargados = False
            error_carga_inicial = error

            traceback.print_exc()

            root.after(
                0,
                lambda:
                lbl_estado.config(
                    text="🔴 Sin conexión"
                ),
            )

            root.after(
                0,
                refrescar_pantalla_seleccion,
            )

    threading.Thread(
        target=iniciar,
        daemon=True,
    ).start()

    estado_archivos_word = (
        obtener_estado_archivos_word()
    )

    threading.Thread(
        target=monitor_word,
        daemon=True,
    ).start()

    threading.Thread(
        target=monitor_sincronizacion,
        daemon=True,
    ).start()

    root.after(
        1500,
        lambda:
        ejecutar_importacion_word(False),
    )

    root.mainloop()


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":
    crear_interfaz()
