"""Utilidades para modificar tablas de inventario en archivos Word."""

from copy import deepcopy
from pathlib import Path


def limpiar_texto(valor):
    return "" if valor is None else str(valor).strip()


def formatear_numero(valor):
    try:
        return f"{float(valor or 0):g}"
    except Exception:
        return str(valor or 0)


# ============================================================
# AGREGAR FILA AL WORD
# ============================================================

def _agregar_fila_word(documento, datos):

    try:
        from docx import Document

    except Exception as error:

        raise Exception(
            "Falta instalar python-docx. "
            "Ejecutá: pip install python-docx"
        ) from error

    ruta = Path(
        str(documento.get("ruta") or "")
    )

    if not ruta.exists():

        raise Exception(
            f"No se encontró el archivo Word:\n{ruta}"
        )

    doc = Document(str(ruta))

    tabla_objetivo = None
    mapa = None

    for tabla_word in doc.tables:

        if not tabla_word.rows:
            continue

        encabezados = [
            limpiar_texto(c.text).lower()
            for c in tabla_word.rows[0].cells
        ]

        mapa_tmp = {}

        for i, encabezado in enumerate(encabezados):

            if encabezado in (
                "codigo",
                "código",
            ):
                mapa_tmp["codigo"] = i

            elif encabezado == "material":
                mapa_tmp["material"] = i

            elif encabezado in (
                "cantidad",
                "cant",
                "cant.",
                "stock",
                "existencia",
                "existencias",
            ):
                mapa_tmp["cantidad"] = i

            elif encabezado == "unidad":
                mapa_tmp["unidad"] = i

            elif encabezado in (
                "categoria",
                "categoría",
            ):
                mapa_tmp["categoria"] = i

            elif encabezado in (
                "ubicacion",
                "ubicación",
            ):
                mapa_tmp["ubicacion"] = i

            elif encabezado in (
                "observaciones",
                "obs",
            ):
                mapa_tmp["observaciones"] = i

        if (
            "material" in mapa_tmp
            and "cantidad" in mapa_tmp
        ):

            tabla_objetivo = tabla_word
            mapa = mapa_tmp
            break

    if tabla_objetivo is None:

        raise Exception(
            "No encontré en el Word una tabla "
            "con las columnas Material y Cantidad."
        )

    nueva_fila = tabla_objetivo.add_row()

    if len(tabla_objetivo.rows) >= 2:

        try:

            fila_modelo = tabla_objetivo.rows[-2]

            for origen, destino in zip(
                fila_modelo.cells,
                nueva_fila.cells,
            ):

                destino._tc.get_or_add_tcPr()

                for hijo in list(
                    origen._tc.tcPr
                ):

                    destino._tc.tcPr.append(
                        deepcopy(hijo)
                    )

        except Exception:
            pass

    valores = {
        "codigo": datos.get("codigo") or "",
        "material": datos.get("material") or "",
        "cantidad": formatear_numero(
            datos.get("cantidad", 0)
        ),
        "unidad": datos.get("unidad") or "",
        "categoria": datos.get("categoria") or "",
        "ubicacion": datos.get("ubicacion") or "",
        "observaciones": datos.get("observaciones") or "",
    }

    for campo, indice in mapa.items():

        if indice < len(
            nueva_fila.cells
        ):

            nueva_fila.cells[
                indice
            ].text = valores.get(
                campo,
                "",
            )

    doc.save(str(ruta))

    return ruta


# ============================================================
# ELIMINAR FILA DEL WORD
# ============================================================

def _eliminar_fila_word(documento, material):

    try:
        from docx import Document

    except Exception as error:

        raise Exception(
            "Falta instalar python-docx. "
            "Ejecutá: pip install python-docx"
        ) from error

    ruta = Path(
        str(documento.get("ruta") or "")
    )

    if not ruta.exists():

        raise Exception(
            f"No se encontró el archivo Word:\n{ruta}"
        )

    doc = Document(str(ruta))

    codigo = limpiar_texto(
        material.get("codigo")
    )

    nombre = limpiar_texto(
        material.get("material")
    )

    fila_encontrada = False

    for tabla_word in doc.tables:

        if not tabla_word.rows:
            continue

        encabezados = [
            limpiar_texto(c.text).lower()
            for c in tabla_word.rows[0].cells
        ]

        indice_codigo = None
        indice_material = None

        for i, encabezado in enumerate(
            encabezados
        ):

            if encabezado in (
                "codigo",
                "código",
            ):

                indice_codigo = i

            elif encabezado == "material":

                indice_material = i

        if (
            indice_codigo is None
            and indice_material is None
        ):
            continue

        if codigo and indice_codigo is not None:

            for fila in list(
                tabla_word.rows[1:]
            ):

                valores = [
                    limpiar_texto(
                        celda.text
                    )
                    for celda in fila.cells
                ]

                if (
                    indice_codigo < len(valores)
                    and valores[indice_codigo] == codigo
                ):

                    tr = fila._tr

                    tr.getparent().remove(tr)

                    fila_encontrada = True
                    break

        elif (
            not codigo
            and nombre
            and indice_material is not None
        ):

            for fila in list(
                tabla_word.rows[1:]
            ):

                valores = [
                    limpiar_texto(
                        celda.text
                    )
                    for celda in fila.cells
                ]

                if (
                    indice_material < len(valores)
                    and valores[indice_material] == nombre
                ):

                    tr = fila._tr

                    tr.getparent().remove(tr)

                    fila_encontrada = True
                    break

        if fila_encontrada:
            break

    if not fila_encontrada:

        raise Exception(
            "No se encontró el material seleccionado "
            "en el archivo Word."
        )

    doc.save(str(ruta))

    return ruta



